"""환경 점검: 모든 핵심 라이브러리 import + 최소 동작 확인 (Gemini 호출 없음).

SETUP.md §13.1 의 스크립트를 기준으로 하고, 설치된 버전의 API 변경에 맞춰 보강했다.
"""

import pathlib
import sys
import tempfile
import warnings

warnings.filterwarnings("ignore")


def ok(m):
    print(f"[OK] {m}")


tmp = pathlib.Path(tempfile.mkdtemp())

assert sys.version_info[:2] == (3, 11), f"Python 3.11 필요, 현재 {sys.version}"
ok(f"python {sys.version.split()[0]}")

# 0) PYTHONPATH 위생: 다른 파이썬 버전 경로가 섞이면 재현 불가 오류가 난다
import re  # noqa: E402

_mine = f"python{sys.version_info.major}.{sys.version_info.minor}"
_bad = [p for p in sys.path if re.search(r"python3\.\d+", p) and _mine not in p]
if _bad:
    print(f"[!] 경고: 다른 파이썬 버전 경로가 sys.path 에 있습니다 → {_bad}")
    print("    실행 전에 `unset PYTHONPATH` 를 권장합니다.")
else:
    ok("sys.path 위생 (다른 버전 경로 없음)")

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

ok(f"numpy {np.__version__} / pandas {pd.__version__}")

# 1) PDF -> 이미지 (Poppler 불필요)
import pymupdf as fitz  # noqa: E402
from PIL import Image  # noqa: E402

doc = fitz.open()
page = doc.new_page(width=842, height=595)  # 가로형(슬라이드 비율)
page.insert_text((72, 72), "Dummy 8D report", fontsize=20)
pdf = tmp / "t.pdf"
doc.save(pdf)
doc.close()
d = fitz.open(pdf)
pix = d[0].get_pixmap(dpi=150)
png = tmp / "p1.png"
pix.save(png)
w, h = Image.open(png).size
ok(f"PyMuPDF {fitz.VersionBind} render {w}x{h}, text_layer={bool(d[0].get_text().strip())}")

# 2) 평가자 일치도
import krippendorff  # noqa: E402
from statsmodels.stats.inter_rater import aggregate_raters, fleiss_kappa  # noqa: E402

r = np.array([[1, 1, 1, 0, 1], [0, 0, 1, 0, 0], [1, 1, 1, 1, 1], [0, 1, 0, 0, 0]])
k = fleiss_kappa(aggregate_raters(r)[0])
a = krippendorff.alpha(reliability_data=r.T, level_of_measurement="nominal")
ok(f"fleiss_kappa={k:.3f}, krippendorff_alpha={a:.3f}")

# 3) 분류·규칙 모델
from scipy.stats import fisher_exact  # noqa: E402
from sklearn.linear_model import LogisticRegression  # noqa: E402
from sklearn.model_selection import RepeatedStratifiedKFold, cross_val_score  # noqa: E402
from sklearn.svm import LinearSVC  # noqa: E402
from statsmodels.stats.multitest import multipletests  # noqa: E402

rng = np.random.default_rng(0)
X = pd.DataFrame(rng.integers(0, 2, (60, 8)), columns=[f"c{i}" for i in range(8)])
y = ((X.c0 + X.c3 + rng.integers(0, 2, 60)) >= 2).astype(int)
svm = LinearSVC(C=0.5, class_weight="balanced").fit(X, y)
s = cross_val_score(LinearSVC(class_weight="balanced"), X, y,
                    cv=RepeatedStratifiedKFold(n_splits=5, n_repeats=2, random_state=0),
                    scoring="balanced_accuracy")
# sklearn 1.8+ : penalty= 폐기 예고 → l1_ratio=1 로 L1 지정
LogisticRegression(l1_ratio=1, solver="liblinear", class_weight="balanced").fit(X, y)
p = [fisher_exact(pd.crosstab(X[c], y))[1] for c in X]
multipletests(p, method="fdr_bh")
ok(f"LinearSVC cv bal_acc={s.mean():.2f}, top weight={X.columns[np.argmax(svm.coef_[0])]}")

from imodels import RuleFitClassifier  # noqa: E402

rf = RuleFitClassifier(max_rules=10, random_state=0).fit(X.values, y.values, feature_names=list(X.columns))
ok(f"imodels RuleFit rules={len(rf._get_rules())}")

# 4) 출력 문서
from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

dx = Document()
st = dx.styles["Normal"]
st.font.name = "Malgun Gothic"
_rpr = st.element.get_or_add_rPr()
_rf = _rpr.find(qn("w:rFonts"))
if _rf is None:
    _rf = OxmlElement("w:rFonts")
    _rpr.append(_rf)
_rf.set(qn("w:eastAsia"), "맑은 고딕")
dx.add_heading("1. 문제 정의 (D2)", level=1)
dx.add_paragraph("[작성 가이드] 예시")
t = dx.add_table(rows=2, cols=3)
t.style = "Table Grid"
dx.save(tmp / "t.docx")

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

pr = Presentation()
pr.slide_width, pr.slide_height = Inches(13.333), Inches(7.5)
sl = pr.slides.add_slide(pr.slide_layouts[5])
sl.shapes.title.text = "근본원인 분석"
pr.save(tmp / "t.pptx")
import docxtpl  # noqa: E402,F401
import openpyxl  # noqa: E402,F401

ok("python-docx / python-pptx / docxtpl / openpyxl")

# 5) Gemini SDK (호출 없이 타입만)
from google import genai  # noqa: E402
from google.genai import types  # noqa: E402
from pydantic import BaseModel  # noqa: E402


class Probe(BaseModel):
    ok: bool


cfg = types.GenerateContentConfig(temperature=0, response_mime_type="application/json", response_schema=Probe)
part = types.Part.from_bytes(data=png.read_bytes(), mime_type="image/png")
assert cfg.response_schema is Probe and part.inline_data is not None
ok(f"google-genai {getattr(genai, '__version__', '')} types OK")

import dotenv  # noqa: E402,F401
import jinja2  # noqa: E402,F401
import matplotlib  # noqa: E402,F401
import rich  # noqa: E402,F401
import tenacity  # noqa: E402,F401
import tqdm  # noqa: E402,F401
import typer  # noqa: E402,F401
import yaml  # noqa: E402,F401

ok("typer / rich / tenacity / yaml / dotenv / jinja2 / matplotlib / tqdm")

# 6) 프로젝트 패키지 자체
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "src"))
from svmtrial import cache, config, features, labels, modeling, schemas  # noqa: E402,F401

ok("svmtrial 패키지 import OK")

print("\nALL SMOKE TESTS PASSED")
