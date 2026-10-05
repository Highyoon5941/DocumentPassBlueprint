"""TemplateSpec → template.pptx (SETUP.md §8.6-3).

- 16:9
- 표지 슬라이드 → 섹션별로 slides_hint 만큼 슬라이드
- 각 슬라이드: 제목, 플레이스홀더 박스(표/그림/텍스트), 발표자 노트에 작성 가이드와 근거
- 마지막 슬라이드에 체크리스트
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from svmtrial.schemas import Element, TemplateSpec

GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
W, H = 13.333, 7.5


def _blank(prs: Presentation):
    """빈 레이아웃(보통 6번)을 고른다. 없으면 마지막 레이아웃."""
    for i in (6, 5):
        if i < len(prs.slide_layouts):
            return prs.slide_layouts[i]
    return prs.slide_layouts[-1]


def _title(slide, text: str, font: str, size: int = 26):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(W - 1), Inches(0.9))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = font
    return box


def _note(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def _placeholder(slide, left, top, width, height, label: str, font: str, dashed: bool = True):
    from pptx.enum.shapes import MSO_SHAPE

    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                                 Inches(width), Inches(height))
    shp.fill.solid()
    shp.fill.fore_color.rgb = LIGHT
    shp.line.color.rgb = GRAY
    shp.line.width = Pt(1)
    if dashed:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE

        shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    tf = shp.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(12)
    r.font.color.rgb = GRAY
    r.font.name = font
    return shp


def _element_label(el: Element) -> str:
    mark = "☑ 필수" if el.required else "☐ 권장"
    kind = {"table": "표", "figure": "그림", "photo": "사진", "chart": "그래프",
            "signature": "서명란", "checklist": "체크리스트", "text": "텍스트"}.get(el.kind, el.kind)
    cols = f"\n열: {', '.join(el.table_columns)}" if el.table_columns else ""
    return f"[{kind}] {el.title}\n({mark}){cols}"


def render(spec: TemplateSpec, out: Path, font: str = "맑은 고딕") -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    layout = _blank(prs)

    # ---- 표지
    s0 = prs.slides.add_slide(layout)
    _title(s0, f"{spec.group.split('__')[0]} 표준 템플릿", font, size=34)
    box = s0.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(W - 1), Inches(2.0))
    tf = box.text_frame
    tf.word_wrap = True
    for line in [f"대상 그룹: {spec.group}", f"타깃 정의: {spec.target_definition}",
                 f"문서 {spec.n_docs}건 중 Pass {spec.n_target_pass}건",
                 "[제목] / [고객사] / [문서번호] / [작성일] 을 채우세요"]:
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size = Pt(14)
        r.font.color.rgb = GRAY
        r.font.name = font
    _placeholder(s0, 0.5, 5.6, 12.3, 1.3, "작성          |          검토          |          승인", font)
    _note(s0, "이 템플릿은 평가자 Pass 조건 역추적 결과로 만들어졌다. 실제 내용은 작성자가 채운다.")

    # ---- 문서 전체 규칙
    if spec.global_rules:
        sg = prs.slides.add_slide(layout)
        _title(sg, "문서 전체 규칙", font)
        n = max(1, len(spec.global_rules))
        h = min(1.1, 5.4 / n)
        for i, el in enumerate(spec.global_rules):
            _placeholder(sg, 0.6, 1.35 + i * (h + 0.12), 12.1, h, _element_label(el).replace("\n", "  "), font)
        _note(sg, "\n\n".join(f"{el.title}: {el.guidance} [근거 {', '.join(el.evidence_ids)}]"
                              for el in spec.global_rules))

    # ---- 섹션
    for sec in spec.sections:
        n_slides = max(1, int(sec.slides_hint or 1))
        per = max(1, -(-len(sec.elements) // n_slides)) if sec.elements else 1
        chunks = [sec.elements[i:i + per] for i in range(0, len(sec.elements), per)] or [[]]
        for si, chunk in enumerate(chunks, start=1):
            sl = prs.slides.add_slide(layout)
            suffix = f" ({si}/{len(chunks)})" if len(chunks) > 1 else ""
            _title(sl, f"{sec.order}. {sec.title}{suffix}" + ("" if sec.required else " [선택]"), font)
            cap = sl.shapes.add_textbox(Inches(0.5), Inches(1.15), Inches(W - 1), Inches(0.4))
            r = cap.text_frame.paragraphs[0].add_run()
            r.text = f"Pass 문서 존재율 {sec.presence_in_pass:.0%}"
            r.font.size = Pt(11)
            r.font.color.rgb = GRAY
            r.font.name = font
            if not chunk:
                _placeholder(sl, 0.6, 1.7, 12.1, 4.9, "[내용 자리]", font)
            else:
                cols = 2 if len(chunk) > 2 else 1
                rows = -(-len(chunk) // cols)
                bw = (12.1 - (cols - 1) * 0.3) / cols
                bh = (5.0 - (rows - 1) * 0.3) / rows
                for i, el in enumerate(chunk):
                    rr, cc = divmod(i, cols)
                    _placeholder(sl, 0.6 + cc * (bw + 0.3), 1.7 + rr * (bh + 0.3), bw, bh,
                                 _element_label(el), font)
            _note(sl, f"{sec.guidance}\n\n" + "\n\n".join(
                f"- {el.title}: {el.guidance} [근거 {', '.join(el.evidence_ids)}]" for el in chunk))

    # ---- 체크리스트
    sc = prs.slides.add_slide(layout)
    _title(sc, "제출 전 체크리스트", font)
    box = sc.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(12.1), Inches(5.6))
    tf = box.text_frame
    tf.word_wrap = True
    items = sorted(spec.evidence, key=lambda e: (e.verdict != "확정", -abs(e.pass_rate_with - e.pass_rate_without)))
    for ev in items[:14]:
        title = next((el.title for s_ in spec.sections for el in s_.elements if ev.concept_id in el.evidence_ids),
                     ev.concept_id)
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = (f"{'☑' if ev.verdict == '확정' else '☐'} [{ev.verdict}] {title} — "
                  f"있을 때 {ev.pass_rate_with:.0%} vs 없을 때 {ev.pass_rate_without:.0%}")
        r.font.size = Pt(13)
        r.font.name = font
    _note(sc, "\n".join(spec.caveats))

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out
