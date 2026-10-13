"""자율 모드 · 로지스틱 계열 · 목적함수 최적화 · LayoutSpec 렌더 (실제 LLM 호출 없음, R6)."""

from __future__ import annotations

import numpy as np
import pandas as pd
import pytest

from svmtrial.modeling import (
    fit_hierarchical,
    fit_logistic,
    fit_monotone_logistic,
    fit_primary,
)
from svmtrial.objective import Objective, optimize, shortfall
from svmtrial.schemas import LayoutBlock, LayoutPage, LayoutSpec


def _world(n=80, seed=3):
    """c0(강) c1(중) 이 진짜 신호, c2 c3 는 미끼. 평가자 3명, E 격 평가자는 c3 의존."""
    rng = np.random.default_rng(seed)
    cols = ["c0", "c1", "c2", "c3"]
    X = pd.DataFrame(rng.integers(0, 2, (n, 4)).astype(float), columns=cols)
    X.index = [f"D{i:03d}" for i in range(n)]
    X.index.name = "doc_id"
    rows = []
    for d in X.index:
        base = -1.2 + 2.0 * X.loc[d, "c0"] + 1.2 * X.loc[d, "c1"]
        for r, bias in (("r1", 0.0), ("r2", 0.5), ("r3", -0.4)):
            z = base + bias + (1.8 * X.loc[d, "c3"] if r == "r3" else 0.0)
            rows.append({"doc_id": d, "rater": r, "y": float(rng.random() < 1 / (1 + np.exp(-z)))})
    long = pd.DataFrame(rows)
    y = long.groupby("doc_id")["y"].mean().ge(0.5).astype(float)
    return X, y.loc[X.index], long, cols


# ---------------------------------------------------------------- 로지스틱


def test_logistic_recovers_signal_order(settings):
    X, y, _, _ = _world()
    out = fit_logistic(settings, X, y)
    assert out["kind"] == "logistic" and out["fit"] in {"mle", "ridge"}
    w = out["weights"]
    assert w["c0"] > w["c1"] > 0, f"신호 순서가 복원되지 않았다: {w}"
    assert w["c0"] > abs(w["c2"]) and w["c0"] > abs(w["c3"])
    # 오즈비는 exp(계수)
    assert out["odds_ratio"]["c0"] == pytest.approx(float(np.exp(w["c0"])), rel=1e-6)


def test_logistic_is_not_over_penalized(settings):
    """기본값이 과벌점이면 계수가 전부 0 근처로 뭉개진다 (회귀 방지)."""
    X, y, _, _ = _world()
    w = fit_logistic(settings, X, y)["weights"]
    assert abs(w["c0"]) > 0.5, f"계수가 과도하게 축소됐다: {w}"


def test_monotone_constraint_forbids_negative(settings):
    X, y, _, cols = _world()
    out = fit_monotone_logistic(X, y, set(cols))
    assert out["converged"]
    assert all(v >= -1e-9 for v in out["weights"].values()), out["weights"]
    assert out["n_constrained"] == len(cols)


def test_hierarchical_uses_rater_level_rows(settings):
    X, _, long, cols = _world()
    out = fit_hierarchical(settings, long, X, cols, rater_interactions=("c3",))
    assert out["n_obs"] == len(long), "평가자 수준 행을 다 쓰지 않았다"
    assert out["n_obs"] > len(X), "문서 수보다 많은 관측치를 써야 한다"
    assert set(out["raters"]) == {"r1", "r2", "r3"}
    assert out["weights"]["c0"] > out["weights"]["c1"] > 0
    # r3 만 c3 에 반응하도록 생성했다 → 그 상호작용이 가장 커야 한다
    inter = {k: v for k, v in out["rater_interactions"].items()}
    assert inter, "상호작용 항이 만들어지지 않았다"
    strongest = max(inter, key=lambda k: inter[k])
    assert "r3" in strongest, f"평가자별 차이를 못 잡았다: {inter}"


def test_hierarchical_accepts_long_with_extra_columns(settings):
    """long 에 특징 열이 섞여 있어도 접미사 충돌 없이 동작해야 한다 (회귀 방지)."""
    X, _, long, cols = _world()
    merged = long.merge(X.reset_index(), on="doc_id")       # 일부러 특징을 섞는다
    out = fit_hierarchical(settings, merged, X, cols)
    assert out["n_obs"] == len(long)


def test_fit_primary_dispatches_by_method(settings):
    X, y, long, cols = _world()
    settings.analysis.method = "svm"
    assert fit_primary(settings, X, y)["kind"] == "svm"
    settings.analysis.method = "logistic"
    assert fit_primary(settings, X, y)["kind"] == "logistic"
    settings.analysis.monotone = True
    assert fit_primary(settings, X, y)["kind"] == "monotone_logistic"
    settings.analysis.monotone = False
    settings.analysis.method = "hierarchical"
    assert fit_primary(settings, X, y, long=long, concept_cols=cols)["kind"] == "hierarchical"


def test_fit_primary_falls_back_when_hierarchical_impossible(settings):
    X, y, _, cols = _world()
    settings.analysis.method = "hierarchical"
    bad = pd.DataFrame({"doc_id": ["D000"], "rater": ["r1"], "y": [1.0]})
    out = fit_primary(settings, X, y, long=bad, concept_cols=cols)
    assert out["kind"] == "logistic" and "fallback_from" in out


# ---------------------------------------------------------------- 목적함수


def _obj(kind="all_pass_prob"):
    fit = {"weights": {"a": 1.2, "b": 1.0, "c": 0.2, "d": -0.8}, "intercept": -2.5,
           "rater_effect": {"r1": 0.0, "r2": -0.6, "r3": 0.4}, "rater_interactions": {}}
    return Objective(fit, ["a", "b", "c", "d"], {k: 0.0 for k in "abcd"}, kind)


def test_budget_curve_is_monotone_in_useful_elements():
    res = optimize(_obj(), ["a", "b", "c"])
    scores = [r["score"] for r in res["budget_curve"]]
    assert scores == sorted(scores), f"유용한 요소만 쓰면 단조 증가해야 한다: {scores}"
    assert res["budget_curve"][1]["elements"] == ["a"], "1개면 가장 강한 요소"
    assert res["search"] == "exhaustive"


def test_all_pass_prob_is_product_over_raters():
    o = _obj("all_pass_prob")
    import math

    expected = 1.0
    for be in o.rater_effect.values():
        expected *= 1 / (1 + math.exp(-(o.b + o.w["a"] + be)))
    assert o.score(("a",)) == pytest.approx(expected, rel=1e-9)


def test_greedy_not_optimal_is_detected():
    """곱 목적함수에서 탐욕이 최적과 달라질 수 있음을 탐지해야 한다."""
    res = optimize(_obj(), ["a", "b", "c", "d"])
    assert isinstance(res["greedy_is_optimal"], bool)
    # 해로운 요소(d, 계수 음수)는 최적 조합에 없어야 한다
    assert "d" not in res["best"]["elements"]


def test_beam_search_kicks_in_for_many_candidates():
    fit = {"weights": {f"f{i}": 0.3 for i in range(25)}, "intercept": -3.0,
           "rater_effect": {"r1": 0.0}, "rater_interactions": {}}
    o = Objective(fit, [f"f{i}" for i in range(25)], {f"f{i}": 0.0 for i in range(25)}, "prob")
    res = optimize(o, [f"f{i}" for i in range(25)], max_k=3, beam=8)
    assert res["search"].startswith("beam")
    assert len(res["budget_curve"]) == 4          # k = 0..3


def test_shortfall_lists_missing_elements():
    o = _obj("prob")
    x = pd.Series({"a": 0.0, "b": 0.0, "c": 0.0, "d": 0.0})
    r = shortfall(o, x, ["a", "b", "c"], target=0.5)
    assert not r["already_meets"]
    assert r["changes"] and "a" in r["changes"]
    assert r["score_after"] > r["score_before"]


def test_shortfall_short_circuits_when_already_met():
    o = _obj("prob")
    x = pd.Series({"a": 1.0, "b": 1.0, "c": 1.0, "d": 0.0})
    r = shortfall(o, x, ["a", "b", "c"], target=0.01)
    assert r["already_meets"] and r["changes"] == []


# ---------------------------------------------------------------- LayoutSpec 렌더


def _layout():
    return LayoutSpec(pages=[LayoutPage(width=595, height=842, blocks=[
        LayoutBlock(kind="title", rect=[50, 50, 545, 85], text="제목", evidence_ids=["skeleton"]),
        LayoutBlock(kind="guidance", rect=[50, 90, 545, 112], text="[작성 가이드] 안내"),
        LayoutBlock(kind="table", rect=[50, 120, 545, 220], columns=["가", "나", "다"], n_rows=3),
        LayoutBlock(kind="figure_placeholder", rect=[50, 230, 545, 380], text="[그림 자리]"),
        LayoutBlock(kind="signature", rect=[330, 700, 545, 760], columns=["작성", "승인"]),
        LayoutBlock(kind="checklist", rect=[50, 400, 545, 460], items=["항목1", "항목2"]),
    ])])


def test_image_render_produces_pdf_and_png(tmp_path):
    from svmtrial.render_image import render

    out = render(_layout(), tmp_path, "t", dpi=72)
    assert (tmp_path / "t.pdf").exists()
    assert len(out["png"]) == 1
    assert (tmp_path / out["png"][0]).exists()
    assert (tmp_path / "t.pdf").stat().st_size < 3_000_000, "폰트 서브셋이 안 되고 있다"


def test_image_render_is_doc_type_agnostic(tmp_path):
    """렌더러는 kind/rect 만 본다 — 문서 종류 이름이 코드에 없어야 한다."""
    from pathlib import Path as _P

    src = (_P(__file__).resolve().parents[1] / "src/svmtrial/render_image.py").read_text(encoding="utf-8")
    for banned in ("대책서", "요구사양서", "8D", "5Why", "문서번호", "고객사"):
        assert banned not in src, f"렌더러에 문서 종류 가정이 있다: {banned}"


def test_docx_from_same_layout(tmp_path):
    from docx import Document

    from svmtrial.render_docx_layout import render

    f = render(_layout(), tmp_path / "t.docx")
    d = Document(f)
    heads = [p.text for p in d.paragraphs if p.style.name.startswith(("Heading", "Title"))]
    assert "제목" in heads
    hdrs = [[c.text for c in t.rows[0].cells] for t in d.tables]
    assert ["가", "나", "다"] in hdrs
    assert ["작성", "승인"] in hdrs, "서명 칸 이름이 LayoutSpec 에서 와야 한다"


def test_docx_layout_renderer_is_doc_type_agnostic():
    from pathlib import Path as _P

    src = (_P(__file__).resolve().parents[1] / "src/svmtrial/render_docx_layout.py").read_text(encoding="utf-8")
    for banned in ("대책서", "요구사양서", "8D", "문서번호", "고객사"):
        assert banned not in src, f"렌더러에 문서 종류 가정이 있다: {banned}"


# ---------------------------------------------------------------- 자율 모드


def test_autonomous_disables_seed():
    from svmtrial.seeds import seed_for

    assert seed_for("대책서"), "비자율 모드에서는 seed 가 있어야 한다"
    assert seed_for("대책서", autonomous=True) == [], "자율 모드는 seed 를 쓰면 안 된다"


def test_offline_layout_handler_respects_observed_page_size(settings):
    from svmtrial import parts as P
    from svmtrial.offline_backend import OfflineBackend
    from svmtrial.payload import wrap

    data = {"observed": {"page_size": [842.0, 595.0]}, "skeleton": [], "checklist": []}
    payload, _ = OfflineBackend(settings).generate(
        model="m", prompt_id="layout_compose.v1",
        parts=[P.text(wrap(data))], schema=LayoutSpec)
    spec = LayoutSpec.model_validate(payload)
    assert spec.pages[0].width == 842.0 and spec.pages[0].height == 595.0


def test_table_columns_recovered_from_concept_phrase():
    from svmtrial.template_spec import _columns_from_phrase

    assert _columns_from_phrase('x "T: Why1 | Why2 | Why3" y') == ["Why1", "Why2", "Why3"]
    assert _columns_from_phrase('x "그래프 추이" y') is None
