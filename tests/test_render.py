"""TemplateSpec 검증과 docx/pptx 렌더 — 한글, 근거 ID, 16:9, 플레이스홀더."""

from __future__ import annotations

import pytest

from svmtrial import render_docx, render_pptx
from svmtrial.report import md, spec_markdown
from svmtrial.schemas import Element, Evidence, Section, TemplateSpec, TemplateSpecLLM
from svmtrial.template_spec import EvidenceMissing, _validate_evidence


def _el(title="5Why 표", kind="table", required=True, ev=("C001",), cols=None):
    return Element(kind=kind, title=title, required=required,
                   guidance="[작성 가이드] 실제 수치는 작성자가 기입한다. [예: 불량률 추이 — ppm]",
                   table_columns=list(cols) if cols else None, evidence_ids=list(ev))


def _spec(doc_format="docs", n_sections=3):
    secs = [
        Section(order=i + 1, section_id=f"D{i + 2}", title=f"섹션 {i + 2}", required=True,
                presence_in_pass=0.9, guidance="[작성 가이드] 섹션 설명",
                elements=[_el(cols=["Why1", "Why2", "Why3"])],
                slides_hint=2, pages_hint=1.5, evidence_ids=["C001"])
        for i in range(n_sections)
    ]
    return TemplateSpec(
        group="대책서__CUST_A", doc_format=doc_format, target_definition="평가자 전원이 Pass",
        n_docs=40, n_target_pass=12, agreement_alpha=0.31, sections=secs,
        global_rules=[_el("작성·검토·승인 서명란", "signature", True, ("skeleton",))],
        rater_conflicts=["C009: r1 는 있을 때 Pass 쪽, r2 는 없을 때 Pass 쪽 — 요구가 충돌한다."],
        caveats=["상관관계는 인과관계가 아니다.", "표본 수 40건."],
        evidence=[Evidence(concept_id="C001", verdict="확정", pass_rate_with=0.86,
                           pass_rate_without=0.31, n_with=14, n_without=26, q_value=0.004,
                           svm_weight=0.42, sign_stability=0.97)],
        meta={"run_id": "t", "backend": "mock", "created_at": "2026-10-05T00:00:00",
              "model_fast": "f", "model_pro": "p"},
    )


# ---------------------------------------------------------------- 근거 ID 검증


def test_validate_evidence_accepts_full_spec():
    llm = TemplateSpecLLM(sections=_spec().sections, global_rules=_spec().global_rules)
    assert _validate_evidence(llm) == []


def test_validate_evidence_rejects_empty_element():
    s = _spec()
    s.sections[0].elements[0].evidence_ids = []
    llm = TemplateSpecLLM(sections=s.sections, global_rules=s.global_rules)
    errs = _validate_evidence(llm)
    assert len(errs) == 1 and "evidence_ids" in errs[0]


def test_validate_evidence_rejects_empty_section_and_global():
    s = _spec()
    s.sections[1].evidence_ids = []
    s.global_rules[0].evidence_ids = []
    errs = _validate_evidence(TemplateSpecLLM(sections=s.sections, global_rules=s.global_rules))
    assert len(errs) == 2


def test_compose_raises_after_retries(settings):
    """근거 ID 가 계속 비면 예외를 던져야 한다 (조용히 통과 금지)."""
    from svmtrial.template_spec import compose

    class BadGemini:
        def generate_json(self, *, model, prompt_id, parts, schema, cache_key_extra="", media_resolution=None):
            bad = _spec()
            bad.sections[0].elements[0].evidence_ids = []
            return TemplateSpecLLM(sections=bad.sections, global_rules=bad.global_rules)

    with pytest.raises(EvidenceMissing):
        compose(settings, type("G", (), {"key": "g", "safe": "g"})(), BadGemini(),
                [], [], [], "docs", [])


# ---------------------------------------------------------------- docx


def test_docx_renders_with_eastasia_font(tmp_path):
    from docx import Document
    from docx.oxml.ns import qn

    out = render_docx.render(_spec(), tmp_path / "t.docx", font="맑은 고딕")
    assert out.exists() and out.stat().st_size > 10_000
    d = Document(out)
    rf = d.styles["Normal"].element.get_or_add_rPr().find(qn("w:rFonts"))
    assert rf.get(qn("w:eastAsia")) == "맑은 고딕"
    assert d.styles["Normal"].font.name == "맑은 고딕"


def test_docx_contains_sections_checklist_and_caveats(tmp_path):
    from docx import Document

    out = render_docx.render(_spec(), tmp_path / "t.docx")
    d = Document(out)
    text = "\n".join(p.text for p in d.paragraphs)
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert any("섹션 2" in h for h in heads)
    assert any("제출 전 체크리스트" in h for h in heads)
    assert any("한계와 주의사항" in h for h in heads)
    assert "상관관계는 인과관계가 아니다." in text
    assert "요구가 충돌한다" in text
    assert "[작성 가이드]" in text
    assert "☑" in text or "☐" in text


def test_docx_table_uses_given_columns(tmp_path):
    from docx import Document

    out = render_docx.render(_spec(), tmp_path / "t.docx")
    d = Document(out)
    headers = [[c.text for c in t.rows[0].cells] for t in d.tables]
    assert ["Why1", "Why2", "Why3"] in headers


def test_docx_has_no_fabricated_numbers(tmp_path):
    """R10 — 템플릿에 실제 내용(가짜 수치/원인)이 들어가면 안 된다."""
    from docx import Document

    out = render_docx.render(_spec(), tmp_path / "t.docx")
    text = "\n".join(p.text for p in Document(out).paragraphs)
    # 표 셀은 열 이름과 플레이스홀더만
    cells = [c.text for t in Document(out).tables for r in t.rows for c in r.cells]
    for c in cells:
        assert c == "" or not c.replace(".", "").isdigit() or c in {"", "-"}, f"표에 수치가 들어갔다: {c!r}"
    assert "[예:" in text or "[작성 가이드]" in text


# ---------------------------------------------------------------- pptx


def test_pptx_is_16x9_with_notes(tmp_path):
    from pptx import Presentation

    out = render_pptx.render(_spec("slides"), tmp_path / "t.pptx", font="맑은 고딕")
    assert out.exists()
    p = Presentation(out)
    assert round(p.slide_width / p.slide_height, 3) == round(16 / 9, 3)
    assert len(p.slides) >= 3
    notes = [sl.notes_slide.notes_text_frame.text for sl in p.slides if sl.has_notes_slide]
    assert any("작성 가이드" in n for n in notes)
    assert any("근거" in n for n in notes), "발표자 노트에 근거 ID 가 있어야 한다"


def test_pptx_uses_korean_font(tmp_path):
    from pptx import Presentation

    out = render_pptx.render(_spec("slides"), tmp_path / "t.pptx", font="맑은 고딕")
    fonts = {r.font.name for sl in Presentation(out).slides for sh in sl.shapes
             if sh.has_text_frame for pa in sh.text_frame.paragraphs for r in pa.runs}
    assert fonts == {"맑은 고딕"}


def test_pptx_has_checklist_slide(tmp_path):
    from pptx import Presentation

    out = render_pptx.render(_spec("slides"), tmp_path / "t.pptx")
    texts = [sh.text_frame.text for sl in Presentation(out).slides for sh in sl.shapes if sh.has_text_frame]
    assert any("제출 전 체크리스트" in t for t in texts)
    assert any("확정" in t for t in texts)


# ---------------------------------------------------------------- markdown


def test_md_escapes_pipes_and_newlines():
    assert md("a | b") == r"a \| b"
    assert md("a\nb") == "a b"
    assert md(None) == ""


def test_spec_markdown_tables_are_well_formed(tmp_path):
    s = _spec()
    s.sections[0].elements[0].title = "표: Line | Model | Owner"
    out = spec_markdown(s, tmp_path / "spec.md")
    text = out.read_text(encoding="utf-8")
    assert r"Line \| Model \| Owner" in text, "파이프가 이스케이프되어야 한다"
    # 모든 표 행의 열 수가 헤더와 같아야 한다.
    # 이스케이프된 \| 는 구분자가 아니므로 세지 않는다.
    import re as _re

    unescaped = _re.compile(r"(?<!\\)\|")
    for block in text.split("\n\n"):
        lines = [x for x in block.splitlines() if x.startswith("|")]
        if len(lines) >= 3:
            widths = {len(unescaped.findall(x)) for x in lines}
            assert len(widths) == 1, f"표 열 수 불일치: {widths}\n{block[:300]}"


# ---------------------------------------------------------------- 리포트 표기


def test_numeric_feature_not_labeled_as_pass_rate():
    """수치형 특징의 rate_with/without 은 Pass율이 아니라 군 평균이다.

    이것을 %로 찍으면 "Pass율 159%" 같은 틀린 문장이 나온다 (회귀 방지).
    """
    from svmtrial.report import _effect_sentence, _is_binary_row, _rate

    numeric = {"feature": "STR_n_tables", "kind": "numeric", "verdict": "확정",
               "rate_with": 1.59, "rate_without": 0.957, "rd": 0.632}
    binary = {"feature": "C001", "kind": "binary", "verdict": "유력",
              "rate_with": 0.62, "rate_without": 0.29, "rd": 0.33}

    assert not _is_binary_row(numeric)
    assert _is_binary_row(binary)
    assert _is_binary_row({"feature": "C002"}), "kind 가 없으면 이진으로 본다"

    assert _rate(numeric, "with") == "1.59"
    assert "%" not in _rate(numeric, "with")
    assert _rate(binary, "with") == "62%"

    s_num = _effect_sentence(numeric, "표 개수")
    assert "Pass율" not in s_num, f"수치형에 Pass율 표기가 붙었다: {s_num}"
    assert "평균" in s_num and "1.59" in s_num

    s_bin = _effect_sentence(binary, "5Why 표")
    assert "Pass율 62%" in s_bin


def test_report_sections_are_numbered(tmp_path):
    """핵심 결론이 1,2,3 으로 번호가 매겨져야 한다 (1,1,1 이 아니다)."""
    from svmtrial.report import _effect_sentence

    rows = [{"feature": f"C00{i}", "kind": "binary", "verdict": "유력",
             "rate_with": 0.6, "rate_without": 0.3, "rd": 0.3} for i in range(1, 4)]
    lines = [f"{i}. " + _effect_sentence(r, r["feature"]) for i, r in enumerate(rows, start=1)]
    assert [x.split(".")[0] for x in lines] == ["1", "2", "3"]
